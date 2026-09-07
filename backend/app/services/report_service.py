import csv
import io
import json
import os
import statistics
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from sqlalchemy.orm import Session

from app.models import domain  # Ensure all SQLAlchemy models are registered
from app.models.geo import CensusTract, Hospital, HospitalCatchment
from app.models.sdoh import EquityIndex, IndicatorCatalog, SDOHIndicator

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
    KeepTogether,
)
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

REPORT_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "generated_reports")
REPORT_DIR = os.path.abspath(REPORT_DIR)
CHART_DIR = os.path.join(REPORT_DIR, "charts")
os.makedirs(REPORT_DIR, exist_ok=True)
os.makedirs(CHART_DIR, exist_ok=True)


def _fetch_data(
    db: Session, hospital_id: int, catchment_id: int, year: int, domains: list
):
    hospital = db.query(Hospital).get(hospital_id) if hospital_id else None
    catchment = (
        db.query(HospitalCatchment).get(catchment_id) if catchment_id else None
    )
    if not catchment and hospital:
        catchment = hospital.catchments[0] if hospital.catchments else None

    if not catchment:
        # Fallback to first available hospital & catchment if none selected
        catchment = db.query(HospitalCatchment).first()
        if catchment and not hospital:
            hospital = catchment.hospital

    if not catchment:
        return None, None, hospital, None, None

    tract_ids = [m.tract_id for m in catchment.memberships]
    catalog_q = db.query(IndicatorCatalog)
    if domains:
        catalog_q = catalog_q.filter(IndicatorCatalog.domain.in_(domains))
    catalog = catalog_q.all()
    if not catalog:
        catalog = db.query(IndicatorCatalog).all()

    # Query SDOH indicator values
    vals = (
        db.query(SDOHIndicator)
        .filter(
            SDOHIndicator.tract_id.in_(tract_ids),
            SDOHIndicator.year == year,
            SDOHIndicator.catalog_id.in_([c.id for c in catalog]),
        )
        .all()
    )

    # Fallback to latest year if selected year has no data
    if not vals:
        latest_year = db.query(SDOHIndicator.year).order_by(SDOHIndicator.year.desc()).first()
        if latest_year:
            year = latest_year[0]
            vals = (
                db.query(SDOHIndicator)
                .filter(
                    SDOHIndicator.tract_id.in_(tract_ids),
                    SDOHIndicator.year == year,
                    SDOHIndicator.catalog_id.in_([c.id for c in catalog]),
                )
                .all()
            )

    equity_rows = (
        db.query(EquityIndex)
        .filter(EquityIndex.tract_id.in_(tract_ids), EquityIndex.year == year)
        .all()
    )
    if not equity_rows:
        latest_eq_year = db.query(EquityIndex.year).order_by(EquityIndex.year.desc()).first()
        if latest_eq_year:
            equity_rows = (
                db.query(EquityIndex)
                .filter(EquityIndex.tract_id.in_(tract_ids), EquityIndex.year == latest_eq_year[0])
                .all()
            )

    equity_map = {e.tract_id: e for e in equity_rows}
    tract_map = {t.id: t for t in db.query(CensusTract).filter(CensusTract.id.in_(tract_ids)).all()}

    rows = []
    by_domain = {}
    for v in vals:
        item = v.catalog_item
        tract = tract_map.get(v.tract_id)
        eq = equity_map.get(v.tract_id)
        rows.append(
            {
                "tract_id": v.tract_id,
                "tract_geoid": tract.geoid if tract else f"Tract {v.tract_id}",
                "population": tract.total_population if tract else None,
                "indicator_code": item.code if item else None,
                "indicator_name": item.name if item else (item.code if item else "Desconocido"),
                "domain": item.domain if item else "General",
                "value": v.value,
                "equity_index": eq.value if eq else None,
                "risk_level": eq.risk_level if eq else "Desconocido",
                "percentile": eq.percentile if eq else None,
            }
        )
        dom = item.domain if item else "General"
        by_domain.setdefault(dom, []).append(v.value)

    return rows, by_domain, hospital, equity_rows, catchment


def _make_bar_chart(by_domain: dict, title: str) -> str:
    path = os.path.join(CHART_DIR, f"chart_{datetime.now().timestamp() * 1000:.0f}.png")
    domains = list(by_domain.keys())
    means = [
        statistics.mean(v) if v else 0 for v in (by_domain.get(d, []) for d in domains)
    ]
    plt.figure(figsize=(9, 4.2), facecolor="#ffffff")
    bars = plt.bar(domains, means, color="#2563eb", edgecolor="#1d4ed8", width=0.55)
    plt.title(title, fontsize=12, fontweight="bold", pad=12, color="#1e293b")
    plt.ylabel("Puntuación Promedio", fontsize=10, color="#475569")
    plt.xticks(rotation=25, ha="right", fontsize=9, color="#334155")
    plt.yticks(fontsize=9, color="#334155")
    plt.grid(axis="y", linestyle="--", alpha=0.3, color="#94a3b8")
    for bar in bars:
        height = bar.get_height()
        plt.text(
            bar.get_x() + bar.get_width() / 2.0,
            height + 0.1,
            f"{height:.1f}",
            ha="center",
            va="bottom",
            fontsize=8,
            fontweight="bold",
            color="#1e293b",
        )
    plt.tight_layout()
    plt.savefig(path, dpi=120)
    plt.close()
    return path


# ---------------------------------------------------------------- PDF Generator
def generate_pdf(db: Session, rows, by_domain, hospital, catchment, title, year):
    filename = f"SDOH_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    path = os.path.join(REPORT_DIR, filename)

    doc = SimpleDocTemplate(
        path,
        pagesize=landscape(A4),
        leftMargin=0.5 * inch,
        rightMargin=0.5 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0f172a"),
        alignment=0,
        spaceAfter=6,
    )
    sub_style = ParagraphStyle(
        "DocSub",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#475569"),
        spaceAfter=12,
    )
    section_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#1e40af"),
        spaceBefore=10,
        spaceAfter=8,
    )

    story = []
    story.append(Paragraph(title, title_style))
    story.append(
        Paragraph(
            f"<b>Centro Médico:</b> {hospital.name if hospital else 'Red General'} &nbsp;|&nbsp; "
            f"<b>Área de Captación:</b> {catchment.name if catchment else 'General'} &nbsp;|&nbsp; "
            f"<b>Año Evaluado:</b> {year} &nbsp;|&nbsp; <b>Fecha de Emisión:</b> {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            sub_style,
        )
    )
    story.append(Spacer(1, 0.1 * inch))

    # Meta KPI Box
    total_pop = sum(set(r["population"] for r in (rows or []) if r["population"]))
    unique_tracts = len(set(r["tract_geoid"] for r in (rows or [])))
    kpi_data = [
        ["Total Census Tracts", "Población Cobertura", "Indicadores Evaluados", "Dominios SDOH"],
        [str(unique_tracts), f"{total_pop:,}", str(len(rows or [])), str(len(by_domain.keys()))],
    ]
    kpi_tbl = Table(kpi_data, colWidths=[2.2 * inch, 2.2 * inch, 2.2 * inch, 2.2 * inch])
    kpi_tbl.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#475569")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#e0e7ff")),
            ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor("#1e3a8a")),
            ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 1), (-1, 1), 12),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#94a3b8")),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
        ])
    )
    story.append(kpi_tbl)
    story.append(Spacer(1, 0.2 * inch))

    if by_domain:
        story.append(Paragraph("Resumen de Determinantes Sociales por Dominio", section_style))
        chart_path = _make_bar_chart(by_domain, "Puntuación Media por Dominio SDOH")
        story.append(Image(chart_path, width=7.8 * inch, height=3.3 * inch))
        story.append(Spacer(1, 0.15 * inch))

        # Domain summary table
        table_data = [["Dominio Temático", "Muestras Registradas", "Mínimo", "Máximo", "Puntuación Promedio"]]
        for dom, values in by_domain.items():
            avg = sum(values) / len(values) if values else 0
            min_v = min(values) if values else 0
            max_v = max(values) if values else 0
            table_data.append([str(dom), str(len(values)), f"{min_v:.2f}", f"{max_v:.2f}", f"{avg:.3f}"])

        tbl = Table(table_data, colWidths=[2.5 * inch, 1.8 * inch, 1.3 * inch, 1.3 * inch, 1.9 * inch])
        tbl.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 9),
                ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#f8fafc"), colors.white]),
                ("FONTSIZE", (0, 1), (-1, -1), 8.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
            ])
        )
        story.append(tbl)

    doc.build(story)
    return filename, path


# ---------------------------------------------------------------- Word Generator
def generate_word(db: Session, rows, by_domain, hospital, catchment, title, year):
    filename = f"SDOH_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    path = os.path.join(REPORT_DIR, filename)
    doc = Document()

    # Title
    heading = doc.add_heading(title, level=0)
    heading.runs[0].font.color.rgb = RGBColor(30, 58, 138)

    doc.add_paragraph(
        f"Centro Médico: {hospital.name if hospital else 'Red General'}\n"
        f"Área de Influencia: {catchment.name if catchment else 'General'}\n"
        f"Año Evaluado: {year}\n"
        f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}\n"
    )

    if by_domain:
        doc.add_heading("1. Resumen por Dominio de Determinantes Sociales", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Medium Shading 1 Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Dominio"
        hdr[1].text = "Indicadores"
        hdr[2].text = "Mínimo"
        hdr[3].text = "Promedio"

        for dom, values in by_domain.items():
            cells = table.add_row().cells
            cells[0].text = str(dom)
            cells[1].text = str(len(values))
            cells[2].text = f"{min(values):.2f}" if values else "0"
            cells[3].text = f"{(sum(values) / len(values)):.3f}" if values else "0"

        doc.add_paragraph("")
        chart_path = _make_bar_chart(by_domain, "Distribución de Puntuaciones por Dominio")
        doc.add_picture(chart_path, width=Inches(6.2))

    if rows:
        doc.add_heading("2. Detalle de Muestras por Census Tract", level=1)
        d_table = doc.add_table(rows=1, cols=6)
        d_table.style = "Light Shading Accent 1"
        d_hdr = d_table.rows[0].cells
        d_hdr[0].text = "Tract GEOID"
        d_hdr[1].text = "Población"
        d_hdr[2].text = "Dominio"
        d_hdr[3].text = "Indicador"
        d_hdr[4].text = "Valor"
        d_hdr[5].text = "Nivel de Riesgo"

        # Add top sample rows
        for r in rows[:120]:
            row_cells = d_table.add_row().cells
            row_cells[0].text = str(r.get("tract_geoid") or "—")
            row_cells[1].text = str(r.get("population") or "—")
            row_cells[2].text = str(r.get("domain") or "—")
            row_cells[3].text = str(r.get("indicator_name") or r.get("indicator_code") or "—")
            row_cells[4].text = f"{r.get('value'):.2f}" if r.get("value") is not None else "—"
            row_cells[5].text = str(r.get("risk_level") or "—")

    doc.save(path)
    return filename, path


# ---------------------------------------------------------------- Excel Generator
def generate_excel(db: Session, rows, by_domain, hospital, catchment, title, year):
    filename = f"SDOH_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    path = os.path.join(REPORT_DIR, filename)
    wb = Workbook()

    # Style definitions
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")

    # Sheet 1: Resumen por Dominio
    ws = wb.active
    ws.title = "Resumen por Dominio"
    ws.append([title])
    ws.cell(row=1, column=1).font = title_font
    ws.append([f"Hospital: {hospital.name if hospital else 'General'}", f"Área: {catchment.name if catchment else 'General'}", f"Año: {year}"])
    ws.append([])

    headers = ["Dominio SDOH", "Cantidad de Muestras", "Valor Mínimo", "Valor Máximo", "Puntuación Promedio"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for dom, values in (by_domain or {}).items():
        avg = sum(values) / len(values) if values else 0
        min_v = min(values) if values else 0
        max_v = max(values) if values else 0
        ws.append([dom, len(values), round(min_v, 2), round(max_v, 2), round(avg, 4)])

    # Sheet 2: Detalle Completo de Indicadores
    ws2 = wb.create_sheet("Detalle Indicadores")
    headers2 = ["Tract GEOID", "Población", "Dominio", "Código Indicador", "Nombre Indicador", "Valor Observado", "Nivel de Riesgo", "Índice de Equidad"]
    ws2.append(headers2)
    for col in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r in (rows or []):
        ws2.append([
            r.get("tract_geoid"),
            r.get("population"),
            r.get("domain"),
            r.get("indicator_code"),
            r.get("indicator_name"),
            r.get("value"),
            r.get("risk_level"),
            r.get("equity_index"),
        ])

    # Adjust column widths
    for sheet in [ws, ws2]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(path)
    return filename, path


# ---------------------------------------------------------------- CSV Generator
def generate_csv(db: Session, rows, by_domain, hospital, catchment, title, year):
    filename = f"SDOH_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    path = os.path.join(REPORT_DIR, filename)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([
            "Hospital",
            "Area_Captacion",
            "Ano",
            "Tract_GEOID",
            "Poblacion",
            "Dominio",
            "Codigo_Indicador",
            "Nombre_Indicador",
            "Valor_Observado",
            "Nivel_Riesgo",
            "Indice_Equidad",
        ])
        h_name = hospital.name if hospital else "General"
        c_name = catchment.name if catchment else "General"
        for r in (rows or []):
            writer.writerow([
                h_name,
                c_name,
                year,
                r.get("tract_geoid"),
                r.get("population"),
                r.get("domain"),
                r.get("indicator_code"),
                r.get("indicator_name"),
                r.get("value"),
                r.get("risk_level"),
                r.get("equity_index"),
            ])
    return filename, path


GENERATORS = {
    "pdf": generate_pdf,
    "word": generate_word,
    "excel": generate_excel,
    "csv": generate_csv,
}
